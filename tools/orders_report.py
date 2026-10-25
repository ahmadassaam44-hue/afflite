#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
تقرير الطلبات: يحول orders.json (نسخة Supabase أو أي نسخة محلية) لملف Excel مرتّب + CSV.

الاستخدام:
  python3 tools/orders_report.py orders.json
  python3 tools/orders_report.py orders.json --out /path/to/folder
  pip install openpyxl          # مطلوب لملف Excel (CSV بيشتغل من غيره)

بيطلّع:
  orders-report-<تاريخ>.xlsx  → 6 شيتات: كل الطلبات / قيد التنفيذ / ملخص / المحافظات / المسوقين / المنتجات
  orders-all.csv              → CSV بنفس أعمدة الشيت الأول (UTF-8 BOM، يفتح في Excel)

⚠️ الملف فيه بيانات عملاء (أسماء/عناوين/تليفونات): احفظه على جهازك،
   وماترفعوش على الريبو (Public) ولا تبعتوش في الشات.
"""

import argparse
import csv
import datetime
import json
import os
import sys

STATUS_AR = {
    'pending': 'قيد التأكيد',
    'confirmed': 'مؤكد',
    'shipped': 'مشحون',
    'delivered': 'تم التوصيل',
    'no_answer': 'لا يرد',
    'cancelled': 'ملغي',
    'returned': 'مرتجع',
}
ACTIVE = ('pending', 'confirmed', 'shipped', 'delivered', 'no_answer')

MAIN_HEADERS = [
    'م', 'رقم الطلب', 'التاريخ (القاهرة)', 'التاريخ المسجل', 'العميل', 'الهاتف', 'هاتف بديل',
    'المحافظة', 'المنطقة', 'العنوان', 'المنتج', 'اللون', 'المقاس',
    'الكمية', 'سعر القطعة', 'الشحن', 'الخصم', 'الإجمالي',
    'التكلفة', 'عمولة المسوق', 'صافي الربح', 'المسوق', 'المنصة', 'الحالة',
    'سبب الإلغاء', 'ملاحظات العميل', 'ملاحظات الإدارة', 'آخر تحديث حالة',
]

DELIVERY_HEADERS = [
    'م', 'رقم الطلب', 'التاريخ (القاهرة)', 'العميل', 'الهاتف', 'هاتف بديل',
    'المحافظة', 'المنطقة', 'العنوان', 'المنتج', 'الكمية', 'الإجمالي', 'الحالة', 'ملاحظات العميل',
]


def parse_dt(s):
    """يحول '2026-10-16T09:39:04-04:00' لكائن وقت، أو None."""
    if not s:
        return None
    try:
        return datetime.datetime.fromisoformat(s)
    except Exception:
        try:
            return datetime.datetime.strptime(s[:19], '%Y-%m-%dT%H:%M:%S')
        except Exception:
            return None


def to_cairo(dt):
    """الوقت المسجل على السيرفر بـ -04:00 → توقيت القاهرة (+03:00) = +7 ساعات."""
    if not dt:
        return ''
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=datetime.timezone(datetime.timedelta(hours=-4)))
    return (dt.astimezone(datetime.timezone.utc) + datetime.timedelta(hours=3)).strftime('%Y-%m-%d %H:%M')


def num(v, default=0.0):
    try:
        return float(v)
    except Exception:
        return default


def normalize(data):
    if isinstance(data, dict):
        for k in ('orders', 'data', 'items'):
            if isinstance(data.get(k), list):
                data = data[k]
                break
    return [o for o in (data or []) if isinstance(o, dict)]


def build_rows(orders):
    rows = []
    for o in orders:
        dt = parse_dt(o.get('at', ''))
        qty = int(num(o.get('qty'), 1) or 1)
        unit = num(o.get('unit'))
        ship = num(o.get('ship'))
        disc = num(o.get('discount'))
        total = num(o.get('total'))
        base = num(o.get('base_price'))
        earn = num(o.get('m_earn'))
        cost = base * qty
        profit = total - cost - ship - earn
        admin_notes = o.get('admin_notes') or []
        admin_txt = ' | '.join(
            (n.get('note') or '') for n in admin_notes if isinstance(n, dict) and n.get('note')
        ) or (o.get('admin_note') or '')
        rows.append({
            'dt': dt or datetime.datetime.min,
            'رقم الطلب': o.get('order_no', ''),
            'التاريخ (القاهرة)': to_cairo(dt),
            'التاريخ المسجل': (o.get('at') or '')[:19],
            'العميل': (o.get('cname') or '').strip(),
            'الهاتف': o.get('cphone', ''),
            'هاتف بديل': o.get('cphone2', ''),
            'المحافظة': o.get('city', ''),
            'المنطقة': o.get('district', ''),
            'العنوان': o.get('address', ''),
            'المنتج': o.get('pname', ''),
            'اللون': o.get('color', ''),
            'المقاس': o.get('size', ''),
            'الكمية': qty,
            'سعر القطعة': unit,
            'الشحن': ship,
            'الخصم': disc,
            'الإجمالي': total,
            'التكلفة': cost,
            'عمولة المسوق': earn,
            'صافي الربح': profit,
            'المسوق': o.get('ref', ''),
            'المنصة': o.get('plat', ''),
            'الحالة': STATUS_AR.get(o.get('status', ''), o.get('status', '')),
            'سبب الإلغاء': o.get('cancel_reason', ''),
            'ملاحظات العميل': o.get('notes', ''),
            'ملاحظات الإدارة': admin_txt,
            'آخر تحديث حالة': to_cairo(parse_dt(o.get('status_at', ''))),
            '_status_key': o.get('status', ''),
        })
    rows.sort(key=lambda r: r['dt'], reverse=True)
    for i, r in enumerate(rows, 1):
        r['م'] = i
    return rows


def summarize_by(rows, key, label):
    out = {}
    for r in rows:
        k = (r.get(key) or '—').strip() or '—'
        d = out.setdefault(k, {'العدد': 0, 'القطع': 0, 'الإجمالي': 0.0, 'صافي الربح': 0.0})
        d['العدد'] += 1
        d['القطع'] += int(r['الكمية'] or 0)
        d['الإجمالي'] += float(r['الإجمالي'] or 0)
        d['صافي الربح'] += float(r['صافي الربح'] or 0)
    res = []
    for k, v in sorted(out.items(), key=lambda kv: -kv[1]['الإجمالي']):
        res.append({label: k, 'عدد الطلبات': v['العدد'], 'القطع': v['القطع'],
                    'إجمالي المبيعات': round(v['الإجمالي'], 2), 'صافي الربح': round(v['صافي الربح'], 2)})
    return res


def write_csv(path, headers, rows):
    with open(path, 'w', newline='', encoding='utf-8-sig') as fh:
        w = csv.DictWriter(fh, fieldnames=headers, extrasaction='ignore')
        w.writeheader()
        w.writerows(rows)


def write_xlsx(path, rows, active_rows, status_rows, city_rows, mk_rows, prod_rows, kpis):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_font = Font(bold=True, color='FFFFFF')
    head_fill = PatternFill('solid', fgColor='1F4E79')

    def add_sheet(title, headers, data, money_cols=(), freeze='A2'):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
        for r in data:
            ws.append([r.get(h, '') for h in headers])
        for idx, h in enumerate(headers, 1):
            letter = get_column_letter(idx)
            if h in money_cols:
                for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                    for cell in row:
                        cell.number_format = '#,##0.00'
            width = max([len(str(h))] + [len(str(r.get(h, ''))) for r in data[:200]] or [10])
            ws.column_dimensions[letter].width = min(45, max(11, width + 3))
        ws.freeze_panes = freeze
        ws.auto_filter.ref = ws.dimensions
        return ws

    # الشيت الأول: كل الطلبات
    add_sheet('كل الطلبات', MAIN_HEADERS, rows,
              money_cols={'سعر القطعة', 'الشحن', 'الخصم', 'الإجمالي', 'التكلفة', 'عمولة المسوق', 'صافي الربح'},
              freeze='C2')
    # قيد المتابعة (غير ملغي/مرتجع)
    add_sheet('قيد المتابعة', DELIVERY_HEADERS, active_rows, money_cols={'الإجمالي'}, freeze='C2')
    # ملخص
    ws = wb.active
    ws.title = 'الملخص'
    ws.append(['مؤشر', 'القيمة'])
    for c in (1, 2):
        ws.cell(row=1, column=c).font = head_font
        ws.cell(row=1, column=c).fill = head_fill
    ws.cell(row=1, column=3, value='إجمالي المبيعات').font = head_font
    ws.cell(row=1, column=3).fill = head_fill
    for label, count, sales in kpis:
        ws.append([label, count])
        ws.cell(row=ws.max_row, column=3, value=round(sales, 2))
    ws.append([])
    ws.append(['توزيع الحالات', 'عدد الطلبات', 'إجمالي المبيعات'])
    for c in (1, 2, 3):
        ws.cell(row=ws.max_row, column=c).font = head_font
        ws.cell(row=ws.max_row, column=c).fill = head_fill
    for r in status_rows:
        ws.append([r['الحالة'], r['عدد الطلبات'], round(r['إجمالي المبيعات'], 2)])
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 16
    ws.column_dimensions['C'].width = 18
    # باقي الشيتات
    add_sheet('حسب المحافظة', ['المحافظة', 'عدد الطلبات', 'القطع', 'إجمالي المبيعات', 'صافي الربح'],
              city_rows, money_cols={'إجمالي المبيعات', 'صافي الربح'})
    add_sheet('حسب المسوق', ['المسوق', 'عدد الطلبات', 'القطع', 'إجمالي المبيعات', 'صافي الربح'],
              mk_rows, money_cols={'إجمالي المبيعات', 'صافي الربح'})
    add_sheet('حسب المنتج', ['المنتج', 'عدد الطلبات', 'القطع', 'إجمالي المبيعات', 'صافي الربح'],
              prod_rows, money_cols={'إجمالي المبيعات', 'صافي الربح'})
    wb.move_sheet('الملخص', offset=-5)
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description='تقرير Excel/CSV من ملف orders.json')
    ap.add_argument('src', help='مسار ملف orders.json')
    ap.add_argument('--out', help='مجلد المخرجات (الافتراضي: مجلد الملف)', default=None)
    args = ap.parse_args()

    if not os.path.exists(args.src):
        print('❌ الملف مش موجود: %s' % args.src)
        sys.exit(2)
    with open(args.src, 'r', encoding='utf-8') as fh:
        try:
            data = json.load(fh)
        except Exception as e:
            print('❌ الملف مش JSON صالح: %s' % e)
            sys.exit(3)

    orders = normalize(data)
    if not orders:
        print('⚠️  مفيش طلبات في الملف')
        sys.exit(0)

    rows = build_rows(orders)
    active = [r for r in rows if r['_status_key'] in ACTIVE]
    out_dir = args.out or os.path.dirname(os.path.abspath(args.src)) or '.'
    os.makedirs(out_dir, exist_ok=True)
    stamp = datetime.datetime.now().strftime('%Y%m%d-%H%M')

    csv_path = os.path.join(out_dir, 'orders-all.csv')
    write_csv(csv_path, MAIN_HEADERS, rows)

    status_rows = summarize_by(rows, 'الحالة', 'الحالة')
    city_rows = summarize_by(rows, 'المحافظة', 'المحافظة')
    mk_rows = summarize_by(rows, 'المسوق', 'المسوق')
    prod_rows = summarize_by(rows, 'المنتج', 'المنتج')

    total_sales = sum(r['الإجمالي'] for r in rows)
    total_cost = sum(r['التكلفة'] for r in rows)
    total_ship = sum(r['الشحن'] for r in rows)
    total_comm = sum(r['عمولة المسوق'] for r in rows)
    total_disc = sum(r['الخصم'] for r in rows)
    total_profit = sum(r['صافي الربح'] for r in rows)
    kpis = [
        ['إجمالي الطلبات', len(rows), total_sales],
        ['طلبات قيد المتابعة (غير ملغية)', len(active), sum(r['الإجمالي'] for r in active)],
        ['ملغي + مرتجع', len(rows) - len(active), sum(r['الإجمالي'] for r in rows if r not in active)],
        ['تكلفة المنتجات', '', total_cost],
        ['إجمالي الشحن', '', total_ship],
        ['إجمالي عمولات المسوقين', '', total_comm],
        ['إجمالي الخصومات', '', total_disc],
        ['صافي الربح (بعد التكلفة والشحن والعمولة)', '', total_profit],
    ]

    xlsx_path = None
    try:
        xlsx_path = os.path.join(out_dir, 'orders-report-%s.xlsx' % stamp)
        write_xlsx(xlsx_path, rows, active, status_rows, city_rows, mk_rows, prod_rows, kpis)
    except ImportError:
        print('⚠️  openpyxl مش مثبت — نزّله بـ: pip install openpyxl (CSV اتعمل عادي)')

    print('\n📊 ملخص سريع')
    print('   عدد الطلبات: %d (منها %d قيد المتابعة)' % (len(rows), len(active)))
    print('   إجمالي المبيعات: %.2f ج.م' % total_sales)
    print('   تكلفة المنتجات: %.2f ج.م' % total_cost)
    print('   إجمالي الشحن: %.2f ج.م' % total_ship)
    print('   إجمالي عمولات المسوقين: %.2f ج.م' % total_comm)
    print('   إجمالي الخصومات: %.2f ج.م' % total_disc)
    print('   صافي الربح: %.2f ج.م' % total_profit)
    print('   الحالات: %s' % '، '.join('%s=%d' % (r['الحالة'], r['عدد الطلبات']) for r in status_rows))
    print('\n📄 CSV:   %s' % csv_path)
    if xlsx_path:
        print('📄 Excel: %s' % xlsx_path)
    print('\n🔒 الملفات دي فيها بيانات عملاء — احتفظ بيها على جهازك وماترفعهاش على الريبو.')


if __name__ == '__main__':
    main()
